"""Review workbook export. Includes every job row; never an official submission."""

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.persistence.models import Job, JobRow, Prediction, SourceRecord


def export_review_workbook(engine, job_id: str) -> bytes:
    with Session(engine) as session:
        job = session.get(Job, job_id)
        if job is None:
            raise LookupError("JOB_NOT_FOUND")
        rows = session.execute(
            select(JobRow, SourceRecord)
            .join(SourceRecord, SourceRecord.id == JobRow.transaction_id)
            .where(JobRow.job_id == job_id)
            .order_by(SourceRecord.sheet, SourceRecord.physical_row)
        ).all()
        predictions = {
            p.transaction_id: p.payload
            for p in session.scalars(select(Prediction).where(Prediction.job_id == job_id))
        }
        if not rows or any(r.status not in {"completed", "failed"} for r, _ in rows):
            raise ValueError("JOB_NOT_FINISHED")
        if any(source.id not in predictions for _, source in rows):
            raise ValueError("RESULTS_INCOMPLETE")
        workbook = Workbook()
        results = workbook.active
        results.title = "Classification results"

        def append(sheet, values):
            safe_values = []
            for v in values:
                if isinstance(v, (dict, list)):
                    v = json.dumps(v, ensure_ascii=False)
                if isinstance(v, str) and str(v).startswith(('=', '+', '-', '@')):
                    v = "'" + str(v)
                safe_values.append(v)
            sheet.append(safe_values)

        append(
            results,
            [
                "Transaction ID",
                "Source sheet",
                "Excel row",
                "Status",
                "Proposed voucher",
                "Alternative",
                "Review reasons",
                "Explanation",
                "Evidence fields",
                "Missing evidence",
            ],
        )
        source_sheet = workbook.create_sheet("Source evidence")
        append(
            source_sheet, ["Transaction ID", "Source sheet", "Excel row", "Cell", "Type", "Original value"]
        )
        for row, source in rows:
            p = predictions[source.id]
            append(
                results,
                [
                    source.id,
                    source.sheet,
                    source.physical_row,
                    p.get("status", "error"),
                    p.get("proposed_label"),
                    p.get("top_alternative"),
                    p.get("reason_codes", []),
                    p.get("rationale_summary", ""),
                    p.get("evidence_paths", []),
                    p.get("missing_evidence", []),
                ],
            )
            for origin in source.payload.get("sources", []):
                for cell in origin.get("cells", []):
                    append(
                        source_sheet,
                        [
                            source.id,
                            source.sheet,
                            source.physical_row,
                            cell["coordinate"],
                            cell["cell_type"],
                            cell["value"],
                        ],
                    )
        notes = workbook.create_sheet("Read me")
        append(notes, ["Purpose", "Development review output; not an organizer submission."])
        append(
            notes,
            [
                "Review",
                "Accepted means a model proposal passed checks, not a verified gold label. Review and error rows are retained.",
            ],
        )
        append(notes, ["Job", job.id])
        append(notes, ["Harness", job.harness_id])
        for sheet in workbook:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            # Make text inert in a fast single pass
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="2F6654")
            for col in "ABCDEFGHIJ":
                sheet.column_dimensions[col].width = 26
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
