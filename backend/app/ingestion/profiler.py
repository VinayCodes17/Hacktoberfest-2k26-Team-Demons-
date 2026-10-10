"""Safe XLSX profiling with ZIP expansion protection and macro rejection.

Reads workbooks in read-only mode. Never evaluates formulas, follows links,
or modifies the source file. Produces a structured profile of sheets, headers,
row counts, and populated-column statistics.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

# ── Safety limits ──────────────────────────────────────────────────────────

MAX_FILE_SIZE = 50 * 1024**2  # 50 MiB on disk
MAX_EXPANDED_SIZE = 200 * 1024**2  # 200 MiB expanded ZIP
MAX_SHEETS = 50
MAX_ROWS = 50_000
MAX_COLUMNS = 500
ALLOWED_EXTENSIONS = {".xlsx"}

# Columns that must NEVER be used as inference input (answer / target labels).
TARGET_LABEL_COLUMNS: set[str] = {
    "voucher category",
    "voucher_type",
    "target",
    "target_label",
    "label",
    "expected_label",
    "answer",
}


@dataclass
class SheetProfile:
    """Profile of a single worksheet."""

    name: str
    max_row: int
    max_column: int
    headers: list[str]
    data_row_count: int
    populated_columns: dict[str, int]
    has_target_column: bool
    target_column_name: str | None


@dataclass
class WorkbookProfile:
    """Complete profile of a workbook file."""

    filename: str
    source_sha256: str
    file_size_bytes: int
    sheet_profiles: list[SheetProfile]
    detected_format: str | None = None  # "format_a", "format_b", or None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _compute_sha256(path: Path) -> str:
    """Compute SHA-256 of file contents."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _safe_header(cell_value: object) -> str:
    """Convert a header cell value to a clean string."""
    if cell_value is None:
        return ""
    return str(cell_value).strip()


def validate_file(path: Path) -> list[str]:
    """Run safety checks on the file before opening. Returns list of errors."""
    errors: list[str] = []
    if not path.is_file():
        errors.append(f"File does not exist: {path.name}")
        return errors
    if path.suffix.lower() not in ALLOWED_EXTENSIONS:
        errors.append(f"Unsupported file extension: {path.suffix}")
        return errors
    size = path.stat().st_size
    if size > MAX_FILE_SIZE:
        errors.append(f"File exceeds {MAX_FILE_SIZE // (1024**2)} MiB limit: {size} bytes")
        return errors
    if size == 0:
        errors.append("File is empty")
        return errors
    try:
        with ZipFile(path) as archive:
            expanded = sum(entry.file_size for entry in archive.infolist())
            if expanded > MAX_EXPANDED_SIZE:
                errors.append(f"Expanded workbook exceeds {MAX_EXPANDED_SIZE // (1024**2)} MiB limit")
            if any("vbaproject" in entry.filename.lower() for entry in archive.infolist()):
                errors.append("Macro content (VBA) is not allowed")
    except Exception as exc:
        errors.append(f"Cannot read as ZIP archive: {exc}")
    return errors


def _profile_sheet(sheet: Worksheet) -> SheetProfile:
    """Profile a single worksheet without modifying it."""
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)

    max_row = sheet.max_row or 0
    max_column = sheet.max_column or 0

    if max_row > MAX_ROWS:
        raise ValueError(f"Sheet '{sheet.title}' exceeds {MAX_ROWS} row limit: {max_row}")
    if max_column > MAX_COLUMNS:
        raise ValueError(f"Sheet '{sheet.title}' exceeds {MAX_COLUMNS} column limit: {max_column}")

    # Read headers from first row
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        first_row = next(rows_iter)
    except StopIteration:
        return SheetProfile(
            name=sheet.title,
            max_row=0,
            max_column=0,
            headers=[],
            data_row_count=0,
            populated_columns={},
            has_target_column=False,
            target_column_name=None,
        )

    headers = [_safe_header(v) for v in first_row]

    # Detect target/answer columns
    target_col_name: str | None = None
    for h in headers:
        if h.lower().strip() in TARGET_LABEL_COLUMNS:
            target_col_name = h
            break

    # Count data rows and populated cells per column
    populated: dict[str, int] = {h: 0 for h in headers if h}
    data_count = 0
    for row_values in rows_iter:
        if not any(v is not None for v in row_values):
            continue
        data_count += 1
        for idx, value in enumerate(row_values):
            if idx < len(headers) and headers[idx] and value is not None:
                populated[headers[idx]] = populated.get(headers[idx], 0) + 1

    return SheetProfile(
        name=sheet.title,
        max_row=max_row,
        max_column=max_column,
        headers=headers,
        data_row_count=data_count,
        populated_columns=populated,
        has_target_column=target_col_name is not None,
        target_column_name=target_col_name,
    )


def profile_workbook(path: Path) -> WorkbookProfile:
    """Profile an XLSX workbook safely. Returns structured profile.

    Never modifies the source file. Runs ZIP safety checks first.
    """
    errors = validate_file(path)
    if errors:
        return WorkbookProfile(
            filename=path.name,
            source_sha256="",
            file_size_bytes=path.stat().st_size if path.is_file() else 0,
            sheet_profiles=[],
            errors=errors,
        )

    digest = _compute_sha256(path)
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if len(book.sheetnames) > MAX_SHEETS:
            return WorkbookProfile(
                filename=path.name,
                source_sha256=digest,
                file_size_bytes=path.stat().st_size,
                sheet_profiles=[],
                errors=[f"Too many sheets: {len(book.sheetnames)} (limit {MAX_SHEETS})"],
            )

        profiles: list[SheetProfile] = []
        profile_errors: list[str] = []
        warnings: list[str] = []

        for sheet in book:
            try:
                sp = _profile_sheet(sheet)
                profiles.append(sp)
                if sp.has_target_column:
                    warnings.append(
                        f"Sheet '{sp.name}' contains target column '{sp.target_column_name}'. "
                        f"This column will be excluded from classification input."
                    )
            except ValueError as exc:
                profile_errors.append(str(exc))

        return WorkbookProfile(
            filename=path.name,
            source_sha256=digest,
            file_size_bytes=path.stat().st_size,
            sheet_profiles=profiles,
            errors=profile_errors,
            warnings=warnings,
        )
    finally:
        book.close()
