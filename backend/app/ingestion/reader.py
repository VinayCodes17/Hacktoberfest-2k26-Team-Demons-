"""Workbook reader: extract rows with full cell provenance.

Reads an XLSX workbook and produces raw row records preserving:
- Original physical Excel row numbers
- Cell coordinates (e.g., "A2", "B15")
- Cell data types (string, number, boolean, date, formula, null)
- Original values without modification

CRITICAL: The target label column (Voucher Category) is extracted separately
and NEVER included in the transaction input records.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.ingestion.profiler import TARGET_LABEL_COLUMNS, validate_file


@dataclass
class RawCell:
    """A single cell with provenance."""

    coordinate: str  # e.g., "A2"
    column_header: str  # original header name
    cell_type: str  # "s", "n", "b", "d", "f", "e", "inlineStr", "null"
    value: object  # original value, None if empty
    is_formula: bool  # True if cell contained a formula


@dataclass
class RawRow:
    """A single row extracted from the workbook with provenance."""

    sheet_name: str
    physical_row: int  # 1-based Excel row number
    cells: list[RawCell]
    target_label: str | None = None  # Extracted answer column, stored separately


@dataclass
class WorkbookData:
    """Complete extracted data from a workbook."""

    filename: str
    source_sha256: str
    sheet_name: str
    header_row: int
    headers: list[str]
    rows: list[RawRow]
    target_column_name: str | None  # Name of the target/answer column if found
    target_labels: dict[int, str]  # physical_row -> label (for evaluation only)
    excluded_columns: list[str]  # Columns excluded from input
    unmapped_metadata_columns: list[str]  # Additional metadata columns
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _cell_type_code(cell) -> str:
    """Map openpyxl cell data_type to our type codes."""
    dt = getattr(cell, "data_type", None)
    if dt is None or cell.value is None:
        return "null"
    # openpyxl data_type: 's' (string), 'n' (numeric), 'b' (boolean),
    # 'd' (date/datetime), 'f' (formula), 'e' (error)
    if dt in ("s", "n", "b", "d", "f", "e"):
        return dt
    return "s"  # default to string


def read_workbook(
    path: Path,
    sheet_name: str,
    header_row: int = 1,
) -> WorkbookData:
    """Read a workbook sheet and extract all rows with cell provenance.

    Arguments:
        path: Path to the XLSX file.
        sheet_name: Name of the sheet to read.
        header_row: 1-based row number containing headers.

    Returns:
        WorkbookData with all rows and their cell provenance.
        Target label columns are extracted separately and NEVER included
        in the cell data of transaction rows.
    """
    errors = validate_file(path)
    if errors:
        return WorkbookData(
            filename=path.name,
            source_sha256="",
            sheet_name=sheet_name,
            header_row=header_row,
            headers=[],
            rows=[],
            target_column_name=None,
            target_labels={},
            excluded_columns=[],
            unmapped_metadata_columns=[],
            errors=errors,
        )

    # Compute file hash
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    digest = h.hexdigest()

    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name not in book.sheetnames:
            return WorkbookData(
                filename=path.name,
                source_sha256=digest,
                sheet_name=sheet_name,
                header_row=header_row,
                headers=[],
                rows=[],
                target_column_name=None,
                target_labels={},
                excluded_columns=[],
                unmapped_metadata_columns=[],
                errors=[f"Sheet '{sheet_name}' not found. Available: {book.sheetnames}"],
            )

        sheet = book[sheet_name]

        # Read all rows from the sheet
        all_rows = list(sheet.iter_rows())
        if not all_rows or len(all_rows) < header_row:
            return WorkbookData(
                filename=path.name,
                source_sha256=digest,
                sheet_name=sheet_name,
                header_row=header_row,
                headers=[],
                rows=[],
                target_column_name=None,
                target_labels={},
                excluded_columns=[],
                unmapped_metadata_columns=[],
                errors=["Sheet is empty or header row is beyond available data"],
            )

        # Extract headers
        header_cells = all_rows[header_row - 1]
        headers: list[str] = []
        for cell in header_cells:
            val = cell.value
            headers.append(str(val).strip() if val is not None else "")

        # Identify target/answer columns
        target_col_idx: int | None = None
        target_col_name: str | None = None
        excluded_cols: list[str] = []
        excluded_indices: set[int] = set()

        for idx, header_str in enumerate(headers):
            if header_str.lower().strip() in TARGET_LABEL_COLUMNS:
                if target_col_idx is None:
                    target_col_idx = idx
                    target_col_name = header_str
                excluded_cols.append(header_str)
                excluded_indices.add(idx)

        # Extract data rows
        raw_rows: list[RawRow] = []
        target_labels: dict[int, str] = {}
        warnings: list[str] = []

        for offset, row_cells in enumerate(all_rows[header_row:]):
            # Skip completely blank rows
            if not any(cell.value is not None for cell in row_cells):
                continue

            physical_row = header_row + 1 + offset

            # Extract target label separately
            target_label: str | None = None
            if target_col_idx is not None and target_col_idx < len(row_cells):
                label_val = row_cells[target_col_idx].value
                if label_val is not None:
                    target_label = str(label_val).strip()
                    target_labels[physical_row] = target_label

            # Build cell list EXCLUDING target columns
            cells: list[RawCell] = []
            for idx, cell in enumerate(row_cells):
                if idx >= len(headers):
                    break
                if idx in excluded_indices:
                    continue  # NEVER include target columns
                if not headers[idx]:
                    continue  # Skip empty headers

                col_letter = get_column_letter(idx + 1)
                coordinate = f"{col_letter}{physical_row}"
                is_formula = getattr(cell, "data_type", None) == "f"

                cells.append(
                    RawCell(
                        coordinate=coordinate,
                        column_header=headers[idx],
                        cell_type=_cell_type_code(cell),
                        value=cell.value,
                        is_formula=is_formula,
                    )
                )

            raw_rows.append(
                RawRow(
                    sheet_name=sheet_name,
                    physical_row=physical_row,
                    cells=cells,
                    target_label=target_label,
                )
            )

        if target_col_name:
            warnings.append(
                f"Target column '{target_col_name}' found and extracted separately. "
                f"It is EXCLUDED from transaction input to prevent label leakage. "
                f"Labels are stored in target_labels for evaluation only."
            )

        # Filter headers to exclude target columns
        input_headers = [h for idx, h in enumerate(headers) if idx not in excluded_indices and h]

        return WorkbookData(
            filename=path.name,
            source_sha256=digest,
            sheet_name=sheet_name,
            header_row=header_row,
            headers=input_headers,
            rows=raw_rows,
            target_column_name=target_col_name,
            target_labels=target_labels,
            excluded_columns=excluded_cols,
            unmapped_metadata_columns=[],
            warnings=warnings,
        )

    finally:
        book.close()
