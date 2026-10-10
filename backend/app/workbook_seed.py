"""Read-only workbook audit and provenance-preserving seed extraction.

Local CLI for P00, not the P02 public upload boundary. Never evaluates formulas
or follows workbook URLs. Transaction values stay out of the public report.
"""

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from app.contracts import SubmissionContract, Taxonomy, TaxonomyEntry, WorkbookContract


def write_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def inspect_workbook(path: Path, names_authority: str | None = None):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 25 * 1024**2:
        raise ValueError("Expected XLSX at most 25 MiB")
    with ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024**2:
            raise ValueError("Expanded workbook exceeds 100 MiB")
        if any("vbaproject" in entry.filename.lower() for entry in archive.infolist()):
            raise ValueError("Macro content is not allowed")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if len(book.sheetnames) > 30:
            raise ValueError("Too many sheets")
        summaries = []
        for sheet in book:
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 200:
                raise ValueError("Sheet dimensions exceed P00 limits")
            summaries.append(
                {"sheet": sheet.title, "physical_rows": sheet.max_row, "columns": sheet.max_column}
            )
        sheet = book["Voucher Ontology"]
        records = iter(sheet.iter_rows(values_only=True))
        headers = next(records)
        entries = []
        for row_num, values in enumerate(records, start=2):
            row = dict(zip(headers, values))
            # The supplied ontology is one contiguous table followed by a
            # blank separator and a footer. Footer prose is not a category.
            if not any(value is not None for value in values):
                break
            entries.append(
                TaxonomyEntry(
                    name=row["Voucher Category"],
                    family=row["Family"],
                    definition=row["Event Definition (interpreted)"],
                    source_url=row["Source URL"],
                    source_row=row_num,
                    boundary=row["Decision Boundary / Caution"],
                    provenance=row["Current Provenance"],
                )
            )
        taxonomy = Taxonomy(
            revision=f"workbook-{digest[:12]}",
            source_sha256=digest,
            names_confirmed=bool(names_authority),
            names_authority=names_authority,
            entries=entries,
        )
        source = book["Organizer_Ready_Input"]
        rows = iter(source.iter_rows())
        columns = [cell.value for cell in next(rows)]
        if any(not isinstance(c, str) or not c.strip() for c in columns) or len(set(columns)) != len(columns):
            raise ValueError("Invalid/duplicate transaction headers")
        if {"voucher_type", "target", "target_label", "label"} & {c.lower() for c in columns}:
            raise ValueError("Target label column in input")
        counts = {column: 0 for column in columns}
        ids = []
        formula_count = 0
        physical_rows = []
        for cells in rows:
            if not any(cell.value is not None for cell in cells):
                continue
            physical_rows.append(cells[0].row)
            row = dict(zip(columns, cells))
            ids.append(row["transaction_id"].value)
            for column, cell in row.items():
                counts[column] += cell.value is not None
                formula_count += cell.data_type == "f"
        if not physical_rows:
            raise ValueError("No source transactions")
        contract = WorkbookContract(
            source_sha256=digest,
            sheet=source.title,
            header_row=1,
            first_data_row=physical_rows[0],
            last_data_row=physical_rows[-1],
            columns=columns,
            excluded_parallel_views=["Transactions_Input", "Agent_Ready_View", "Agent_Context_Complete"],
        )
        report = {
            "source_file": path.name,
            "source_sha256": digest,
            "sheets": summaries,
            "category_count": len(entries),
            "transaction_count": len(physical_rows),
            "transaction_columns": len(columns),
            "unique_transaction_ids": len(set(ids)),
            "missing_transaction_ids": sum(value is None for value in ids),
            "formula_cells_in_selected_input": formula_count,
            "populated_cells_by_column": counts,
            "data_role": "synthetic_unverified",
            "source_modified": False,
            "selection_reason": "Full source fields; alternative views are not additional transactions",
            "trusted_labels_imported": 0,
        }
        return taxonomy, contract, report
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("../storage/p00"))
    parser.add_argument("--names-authority", default=None)
    args = parser.parse_args()
    taxonomy, contract, report = inspect_workbook(args.workbook, args.names_authority)
    for name, value in [
        ("taxonomy.json", taxonomy.model_dump()),
        ("workbook-contract.json", contract.model_dump()),
        ("workbook-profile.json", report),
        ("submission-contract.json", SubmissionContract().model_dump()),
    ]:
        write_json(args.output / name, value)
    print(
        json.dumps(
            {key: report[key] for key in ["category_count", "transaction_count", "transaction_columns"]}
        )
    )


if __name__ == "__main__":
    main()
