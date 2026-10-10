from openpyxl import Workbook

from app.ingestion.pipeline import ingest_workbook


def test_sparse_alias_columns_and_empty_voucher_are_allowed(tmp_path):
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Transactions"
    sheet.append(["Document Date", "PO Date", "Payment Date", "Narration", "Voucher Category"])
    sheet.append([None, "2026-10-10", None, "Supplier order placed", None])
    sheet.append([None, None, "2026-10-11", "Invoice paid", None])
    book.save(path)
    result = ingest_workbook(path)
    assert result.total_rows == 2
    assert not result.mapping.ambiguous_mappings
    assert "Voucher Category" in result.mapping.excluded_columns
    for index, coordinate in enumerate(["B2", "C3"]):
        signal = next(s for s in result.transactions[index].signals if s.canonical_name == "document_date")
        assert signal.source_coordinate == coordinate
        assert not result.transactions[index].parse_issues
        assert all("voucher" not in s.canonical_name.lower() for s in result.transactions[index].signals)


def test_simultaneous_alias_values_are_preserved_and_flagged(tmp_path):
    path = tmp_path / "input.xlsx"
    book = Workbook()
    book.active.append(["Document Date", "PO Date", "Voucher Category"])
    book.active.append(["2026-10-10", "2026-10-11", "Purchase"])
    book.save(path)
    result = ingest_workbook(path)
    transaction = result.transactions[0]
    assert not result.mapping.ambiguous_mappings
    assert {s.canonical_name for s in transaction.signals} == {
        "document_date[Document Date]",
        "document_date[PO Date]",
    }
    assert len(transaction.document) == 2
    assert transaction.parse_issues[0].startswith("MULTIPLE_SOURCE_VALUES:")
    assert all(s.value != "Purchase" for s in transaction.signals)


def test_duplicate_headers_still_require_correction(tmp_path):
    path = tmp_path / "input.xlsx"
    book = Workbook()
    book.active.append(["Document Date", "Document Date"])
    book.active.append(["2026-10-10", "2026-10-11"])
    book.save(path)
    assert ingest_workbook(path).mapping.ambiguous_mappings
