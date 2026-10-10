from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from test_worker_runtime import queued as queued

from app.persistence.models import SourceRecord
from app.persistence.repository import claim_job_row_lease, record_prediction
from app.review_export import export_review_workbook


def test_export_blocks_incomplete_jobs(queued):
    engine, _, job = queued
    with pytest.raises(ValueError, match="JOB_NOT_FINISHED"):
        export_review_workbook(engine, job.id)
    with pytest.raises(LookupError):
        export_review_workbook(engine, "missing")


def test_export_retains_review_and_inert_source_text(queued):
    engine, _, job = queued
    with Session(engine) as session:
        source = session.get(SourceRecord, "row")
        payload = deepcopy(source.payload)
        payload["sources"][0]["cells"][0]["value"] = '=HYPERLINK("https://example.com")'
        source.payload = payload
        session.commit()
    claim_job_row_lease(engine, "worker")
    record_prediction(
        engine,
        job.id,
        "row",
        1,
        worker_id="worker",
        payload={
            "status": "review",
            "proposed_label": "Payment",
            "reason_codes": ["MISSING_EVIDENCE"],
            "rationale_summary": "=2+2",
            "evidence_paths": ["description"],
        },
    )
    book = load_workbook(BytesIO(export_review_workbook(engine, job.id)), data_only=False)
    sheet = book["Classification results"]
    assert sheet.max_row == 2
    assert sheet["D2"].value == "review"
    assert sheet["H2"].value == "=2+2" and sheet["H2"].data_type == "s"
    cell = book["Source evidence"]["F2"]
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"
    assert book["Read me"]["B2"].value


def test_failed_row_is_included_in_download(queued):
    engine, _, job = queued
    for number in (1, 2):
        claim_job_row_lease(engine, "worker")
        record_prediction(engine, job.id, "row", number, error="TIMEOUT", worker_id="worker")
    book = load_workbook(BytesIO(export_review_workbook(engine, job.id)))
    assert book["Classification results"]["D2"].value == "error"
    assert "TIMEOUT" in book["Classification results"]["G2"].value
